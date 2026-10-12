"""Export a common reading document; no recognition or editorial rewriting."""
from pathlib import Path
from html import escape
import os

from .reading import reading_document


def export_reading(project, path, kind):
    document = reading_document(project)
    if kind == 'txt':
        lines = [document['title'], '']
        lines += [f'{key}：{value}' for key, value in document['metadata']]
        lines += ['']
        for block in document['blocks']:
            lines.extend([block['text'], ''])
        Path(path).write_text('\n'.join(lines), encoding='utf-8-sig')
    elif kind == 'docx':
        _docx(document, path)
    elif kind == 'pdf':
        _pdf(document, path)
    else:
        raise ValueError('不支持的阅读文档格式：'+kind)


def _docx(data, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    for element in (document.styles.element, document.element):
        for border in element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in [('Normal', 11), ('Title', 20), ('Heading 1', 14)]:
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.line_spacing = Pt(28 if name == 'Title' else 21 if name == 'Heading 1' else 17)
        style.paragraph_format.space_after = Pt(10)
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(snap)
    document.add_paragraph(data['title'], 'Title')
    for key, value in data['metadata']:
        p = document.add_paragraph(f'{key}：{value}')
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9)
    for block in data['blocks']:
        document.add_paragraph(block['text'], 'Heading 1' if block['type'] == 'heading' else 'Normal')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    document.save(str(path))


def _pdf(data, path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    font = 'MeetingChinese'
    if font not in pdfmetrics.getRegisteredFontNames():
        font_file = Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts'/'msyh.ttc'
        if font_file.is_file():
            pdfmetrics.registerFont(TTFont(font, str(font_file), subfontIndex=0))
        else:
            font = 'STSong-Light'
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(UnicodeCIDFont(font))
    base = dict(fontName=font, textColor=colors.black, alignment=TA_LEFT, wordWrap='CJK')
    body = ParagraphStyle('Body', fontSize=11, leading=17, spaceAfter=11, **base)
    title = ParagraphStyle('Title', fontSize=20, leading=28, spaceAfter=16, keepWithNext=True, **base)
    heading = ParagraphStyle('Heading', fontSize=14, leading=21, spaceBefore=8, spaceAfter=8, keepWithNext=True, **base)
    meta = ParagraphStyle('Meta', fontSize=9, leading=14, spaceAfter=4, **base)
    story = [Paragraph(escape(data['title']), title)]
    story += [Paragraph(escape(f'{key}：{value}'), meta) for key, value in data['metadata']]
    story.append(Spacer(1, 10))
    story += [Paragraph(escape(block['text']).replace('\n', '<br/>'), heading if block['type']=='heading' else body)
              for block in data['blocks']]
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 9)
        canvas.drawRightString(A4[0]-22*mm, 12*mm, f'第 {doc.page} 页')
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=22*mm, leftMargin=22*mm,
                      topMargin=20*mm, bottomMargin=20*mm, title=data['title']).build(
                          story, onFirstPage=footer, onLaterPages=footer)
